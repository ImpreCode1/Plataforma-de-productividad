from collections import defaultdict
from sqlalchemy.orm import Session
from sqlalchemy import func
from uuid import UUID

from app.models.user import User
from app.models.role import Role, UserRole
from app.models.indicator_assignment import IndicatorAssignment
from app.models.tracking import IndicatorTracking
from app.models.action_plan import ActionPlan
from app.models.evidence import Evidence


def get_dashboard_by_user(db: Session, user_id: UUID, year: int, month: int = None):

    assignments = db.query(IndicatorAssignment).filter(
        IndicatorAssignment.user_id == user_id,
        IndicatorAssignment.year == year,
        IndicatorAssignment.is_active == True
    ).all()

    result = []

    for assignment in assignments:

        trackings = db.query(IndicatorTracking).filter(
            IndicatorTracking.assignment_id == assignment.id
        ).order_by(IndicatorTracking.month).all()

        months = []

        if trackings:
            for t in trackings:
                if month and t.month != month:
                    continue

                action_plans = db.query(ActionPlan).filter(
                    ActionPlan.tracking_id == t.id
                ).all()
                
                plans_data = []
                for plan in action_plans:
                    plans_data.append({
                        "id": str(plan.id),
                        "reason_not_met": plan.reason_not_met,
                        "action_plan": plan.action_plan,
                        "created_at": plan.created_at.isoformat() if plan.created_at else None
                    })

                evidence_count = db.query(Evidence).filter(
                    (Evidence.tracking_id == t.id) |
                    ((Evidence.tracking_id == None) & (Evidence.user_id == t.user_id) & (Evidence.year == t.year) & (Evidence.month == t.month))
                ).count()

                months.append({
                    "month": t.month,
                    "achieved_value": t.achieved_value,
                    "achieved_total": t.achieved_total,
                    "achievement_percentage": t.achievement_percentage,
                    "status": t.status,
                    "is_closed": t.is_closed,
                    "approval_status": t.approval_status or "PENDIENTE",
                    "rejection_comment": t.rejection_comment,
                    "tracking_id": str(t.id),
                    "action_plans": plans_data,
                    "evidence_count": evidence_count
                })
        else:
            month_val = assignment.month
            
            if month and month_val != month:
                continue

            months.append({
                "month": month_val,
                "achieved_value": None,
                "achieved_total": None,
                "achievement_percentage": None,
                "status": "PENDING",
                "is_closed": False,
                "approval_status": "PENDIENTE",
                "rejection_comment": None,
                "tracking_id": None,
                "action_plans": [],
                "evidence_count": 0
            })

        if month and not months:
            continue

        result.append({
            "id": str(assignment.id),
            "indicator_name": assignment.indicator_name,
            "formula": assignment.formula,
            "target_value": assignment.target_value,
            "weight": assignment.weight,
            "months": months
        })

    return {
        "user_id": user_id,
        "year": year,
        "month": month,
        "indicators": result
    }
    
def get_team_dashboard(db: Session, leader_id: UUID, year: int, month: int = None):

    # 1. Obtener subordinados
    team = db.query(User).filter(
        User.leader_id == leader_id,
        User.is_active == True
    ).all()

    result = []

    for user in team:

        # reutilizamos función existente 🔥
        user_dashboard = get_dashboard_by_user(db, user.id, year, month)

        result.append({
            "user_id": user.id,
            "name": user.name,
            "email": user.email,
            "position_name": user.position_name,
            "indicators": user_dashboard["indicators"]
        })

    return {
        "leader_id": leader_id,
        "year": year,
        "month": month,
        "team": result
    }


# ------------------------------------------------
# GLOBAL DASHBOARD (ADMIN) - All users overview
# ------------------------------------------------

def get_global_dashboard(db: Session, year: int, month: int = None, quarter: int = None, area: str = None, direccion: str = None, responsable: str = None, cumplimiento: str = None, search: str = None):
    from app.modules.users.service import normalize_area
    from collections import defaultdict

    all_users = db.query(User).filter(User.is_active == True).all()

    if area:
        normalized_area = normalize_area(area)
        if normalized_area:
            all_users = [
                u for u in all_users
                if normalize_area(u.area) == normalized_area
            ]

    if direccion:
        from app.modules.users.service import normalize_direccion
        normalized_direccion = normalize_direccion(direccion)
        if normalized_direccion:
            all_users = [
                u for u in all_users
                if normalize_direccion(u.direccion) == normalized_direccion
            ]

    if responsable:
        all_users = [u for u in all_users if u.name == responsable]

    if search:
        search_lower = search.lower().strip()
        leader_name_map = {u.id: u.name for u in all_users}

        def matches_search(u):
            if search_lower in (u.name or "").lower():
                return True
            if search_lower in (u.position_name or "").lower():
                return True
            leader_name = leader_name_map.get(u.leader_id, "") if u.leader_id else ""
            if search_lower in (leader_name or "").lower():
                return True
            return False

        all_users = [u for u in all_users if matches_search(u)]

    if not all_users:
        return empty_global_dashboard(year, month)

    filter_month = None
    if month:
        filter_month = [month]
    elif quarter:
        quarter_map = {1: [1, 2, 3], 2: [4, 5, 6], 3: [7, 8, 9], 4: [10, 11, 12]}
        filter_month = quarter_map.get(quarter)

    user_ids = [u.id for u in all_users]
    user_map = {u.id: u for u in all_users}

    total_users = len(all_users)
    total_indicators = 0
    total_tracked = 0
    total_closed = 0
    total_action_plans = 0
    total_evidence = 0

    monthly_stats = {m: {"total": 0, "closed": 0, "plans": 0, "evidence": 0, "score": 0} for m in range(1, 13)}

    # 1. BATCH: all assignments for the year
    all_assignments = db.query(IndicatorAssignment).filter(
        IndicatorAssignment.user_id.in_(user_ids),
        IndicatorAssignment.year == year,
        IndicatorAssignment.is_active == True
    ).all()

    assigments_by_user = defaultdict(list)
    for a in all_assignments:
        assigments_by_user[a.user_id].append(a)

    # 2. BATCH: all trackings for the year
    tracking_query = db.query(IndicatorTracking).filter(
        IndicatorTracking.user_id.in_(user_ids),
        IndicatorTracking.year == year
    )
    if filter_month:
        tracking_query = tracking_query.filter(IndicatorTracking.month.in_(filter_month))
    all_trackings = tracking_query.all()

    trackings_by_user = defaultdict(list)
    tracking_ids = []
    for t in all_trackings:
        trackings_by_user[t.user_id].append(t)
        tracking_ids.append(t.id)

    # 3. BATCH: action plans
    plans_by_tracking = defaultdict(int)
    plans_detail_by_tracking = defaultdict(list)
    if tracking_ids:
        plan_counts = db.query(
            ActionPlan.tracking_id, func.count(ActionPlan.id)
        ).filter(
            ActionPlan.tracking_id.in_(tracking_ids)
        ).group_by(ActionPlan.tracking_id).all()
        for tid, cnt in plan_counts:
            plans_by_tracking[tid] = cnt

        plans = db.query(ActionPlan).filter(
            ActionPlan.tracking_id.in_(tracking_ids)
        ).all()
        for p in plans:
            plans_detail_by_tracking[p.tracking_id].append({
                "id": str(p.id),
                "reason_not_met": p.reason_not_met,
                "action_plan": p.action_plan,
                "created_at": p.created_at.isoformat() if p.created_at else None,
            })

    # 4. BATCH: evidences linked to trackings
    ev_by_tracking = defaultdict(int)
    if tracking_ids:
        ev_counts = db.query(
            Evidence.tracking_id, func.count(Evidence.id)
        ).filter(
            Evidence.tracking_id.in_(tracking_ids)
        ).group_by(Evidence.tracking_id).all()
        for tid, cnt in ev_counts:
            ev_by_tracking[tid] = cnt

    # 5. BATCH: evidences without tracking per user
    ev_without_q = db.query(
        Evidence.user_id, func.count(Evidence.id)
    ).filter(
        Evidence.user_id.in_(user_ids),
        Evidence.year == year,
        Evidence.tracking_id == None
    )
    if filter_month:
        ev_without_q = ev_without_q.filter(Evidence.month.in_(filter_month))
    ev_without_by_user = defaultdict(int, ev_without_q.group_by(Evidence.user_id).all())

    # 6. BATCH: leader names
    leader_ids = set(u.leader_id for u in all_users if u.leader_id)
    leaders = {}
    if leader_ids:
        leader_users = db.query(User).filter(User.id.in_(leader_ids)).all()
        leaders = {l.id: l.name for l in leader_users}

    # 7. BATCH: subordinate counts (is_leader check)
    sub_counts = db.query(
        User.leader_id, func.count(User.id)
    ).filter(
        User.leader_id.in_(user_ids)
    ).group_by(User.leader_id).all()
    is_leader_map = {lid: True for lid, _ in sub_counts}

    # 8. BATCH: user roles
    user_roles_data = db.query(UserRole).filter(
        UserRole.user_id.in_(user_ids)
    ).all()
    role_ids = set(ur.role_id for ur in user_roles_data)
    role_name_map = {}
    if role_ids:
        roles = db.query(Role).filter(Role.id.in_(role_ids)).all()
        role_name_map = {r.id: r.name for r in roles}
    role_names_by_user = defaultdict(set)
    for ur in user_roles_data:
        if ur.role_id in role_name_map:
            role_names_by_user[ur.user_id].add(role_name_map[ur.role_id])

    # 9. BATCH: evidences per user+tracking for indicator detail
    ev_detail_q = db.query(
        Evidence.tracking_id, Evidence.user_id, Evidence.year, Evidence.month, func.count(Evidence.id)
    ).filter(
        Evidence.user_id.in_(user_ids),
        Evidence.year == year,
    )
    if filter_month:
        ev_detail_q = ev_detail_q.filter(Evidence.month.in_(filter_month))
    ev_detail_counts = ev_detail_q.group_by(
        Evidence.tracking_id, Evidence.user_id, Evidence.year, Evidence.month
    ).all()
    ev_detail_map = defaultdict(int)
    for tid, uid, y, m, cnt in ev_detail_counts:
        ev_detail_map[(tid, uid, y, m)] = cnt

    team_summary = []

    for user in all_users:
        uid = user.id
        is_leader = uid in is_leader_map
        roles_for_user = role_names_by_user.get(uid, set())
        is_admin_role = "ADMIN" in roles_for_user
        is_leader_role = "LEADER" in roles_for_user

        raw_assignments = assigments_by_user.get(uid, [])
        if filter_month:
            assignments = [a for a in raw_assignments if a.month in filter_month]
        else:
            assignments = raw_assignments
        user_indicators = len(assignments)
        total_indicators += user_indicators

        trackings = trackings_by_user.get(uid, [])

        if cumplimiento:
            if cumplimiento == "cumplio":
                trackings = [t for t in trackings if t.target_met == True]
            elif cumplimiento == "no_cumplio":
                trackings = [t for t in trackings if t.target_met == False]
        tracked_count = len([t for t in trackings if t.status in ["COMPLETED", "CLOSED"]])
        closed_count = len([t for t in trackings if t.is_closed])
        total_tracked += tracked_count
        total_closed += closed_count

        user_plans = 0
        evidence_from_trackings = 0
        for t in trackings:
            p_cnt = plans_by_tracking.get(t.id, 0)
            user_plans += p_cnt
            e_cnt = ev_by_tracking.get(t.id, 0)
            evidence_from_trackings += e_cnt

            ms = monthly_stats[t.month]
            ms["total"] += 1
            if t.is_closed:
                ms["closed"] += 1
            ms["plans"] += p_cnt
            ms["evidence"] += 1
            if t.weighted_score:
                ms["score"] += t.weighted_score

        evidence_without = ev_without_by_user.get(uid, 0)
        user_evidence = evidence_from_trackings + evidence_without
        total_action_plans += user_plans
        total_evidence += user_evidence

        leader_name = leaders.get(user.leader_id) if user.leader_id else "Sin líder"

        user_score = 0
        if trackings and assignments:
            indicator_groups = {}
            for assignment in assignments:
                ind_name = assignment.indicator_name
                if ind_name not in indicator_groups:
                    indicator_groups[ind_name] = {"weight": assignment.weight or 0, "trackings": []}
                indicator_groups[ind_name]["trackings"].extend(
                    [t for t in trackings if t.assignment_id == assignment.id]
                )

            for ind_data in indicator_groups.values():
                trackings_with_data = [t for t in ind_data["trackings"] if t.achievement_percentage is not None]
                if trackings_with_data:
                    avg = sum(t.achievement_percentage for t in trackings_with_data) / len(trackings_with_data)
                    user_score += (avg * ind_data["weight"]) / 100

        indicators_by_name = {}
        for assignment in assignments:
            trackings_ind = [t for t in trackings if t.assignment_id == assignment.id]

            if trackings_ind:
                for t in trackings_ind:
                    if filter_month and t.month not in filter_month:
                        continue

                    plans_count = plans_by_tracking.get(t.id, 0)
                    evidence_count = ev_by_tracking.get(t.id, 0) + ev_detail_map.get((None, t.user_id, t.year, t.month), 0)

                    month_data = {
                        "month": t.month,
                        "tracking_id": str(t.id),
                        "is_closed": t.is_closed,
                        "achieved_value": float(t.achieved_value) if t.achieved_value else None,
                        "achieved_total": float(t.achieved_total) if t.achieved_total else None,
                        "achievement_percentage": float(t.achievement_percentage) if t.achievement_percentage else None,
                        "weighted_score": float(t.weighted_score) if t.weighted_score else None,
                        "target_met": t.target_met,
                        "plans_count": plans_count,
                        "evidence_count": evidence_count,
                        "has_action_plan": plans_count > 0,
                        "has_evidence": evidence_count > 0,
                        "action_plans": plans_detail_by_tracking.get(t.id, []),
                    }

                    if assignment.indicator_name not in indicators_by_name:
                        indicators_by_name[assignment.indicator_name] = {
                            "indicator_name": assignment.indicator_name,
                            "weight": assignment.weight,
                            "target_value": assignment.target_value,
                            "months": []
                        }
                    indicators_by_name[assignment.indicator_name]["months"].append(month_data)
            else:
                if filter_month and assignment.month not in filter_month:
                    continue

                month_data = {
                    "month": assignment.month,
                    "tracking_id": None,
                    "is_closed": False,
                    "achieved_value": None,
                    "achieved_total": None,
                    "achievement_percentage": None,
                    "weighted_score": None,
                    "target_met": None,
                    "plans_count": 0,
                    "evidence_count": 0,
                    "has_action_plan": False,
                    "has_evidence": False,
                    "action_plans": [],
                }

                if assignment.indicator_name not in indicators_by_name:
                    indicators_by_name[assignment.indicator_name] = {
                        "indicator_name": assignment.indicator_name,
                        "weight": assignment.weight,
                        "target_value": assignment.target_value,
                        "months": []
                    }
                indicators_by_name[assignment.indicator_name]["months"].append(month_data)

        indicators_detail = list(indicators_by_name.values())

        indicators_count_filtered = len(indicators_detail) if cumplimiento else None

        if cumplimiento and not indicators_detail:
            continue

        team_summary.append({
            "user_id": str(uid),
            "name": user.name,
            "email": user.email,
            "position_name": user.position_name,
            "leader_name": leader_name,
            "indicators_count": user_indicators,
            "indicators_count_filtered": indicators_count_filtered,
            "tracked_months": tracked_count,
            "closed_months": closed_count,
            "score": round(user_score, 2),
            "action_plans": user_plans,
            "evidence_count": user_evidence,
            "indicators": indicators_detail,
            "is_leader": is_leader,
            "is_leader_role": is_leader_role,
            "is_admin_role": is_admin_role,
        })

    team_summary.sort(key=lambda x: x["score"], reverse=True)

    teams_data = {}
    for user_entry in team_summary:
        leader = user_entry["leader_name"]
        if leader not in teams_data:
            teams_data[leader] = {
                "leader_name": leader,
                "members": [],
                "total_indicators": 0,
                "total_indicators_filtered": 0,
                "total_closed": 0,
                "total_plans": 0,
                "total_evidence": 0,
                "total_score": 0,
            }
        teams_data[leader]["members"].append(user_entry)
        teams_data[leader]["total_indicators"] += user_entry["indicators_count"]
        if user_entry.get("indicators_count_filtered") is not None:
            teams_data[leader]["total_indicators_filtered"] += user_entry["indicators_count_filtered"]
        teams_data[leader]["total_closed"] += user_entry["closed_months"]
        teams_data[leader]["total_plans"] += user_entry["action_plans"]
        teams_data[leader]["total_evidence"] += user_entry["evidence_count"]
        teams_data[leader]["total_score"] += user_entry["score"]

    for team in teams_data.values():
        members_with_data = [m for m in team["members"] if m["indicators_count"] > 0]
        member_count = len(members_with_data) if members_with_data else 1
        team["avg_score"] = round(team["total_score"] / member_count, 2)
        team["members"].sort(key=lambda x: x["score"], reverse=True)

    teams_list = sorted(teams_data.values(), key=lambda x: x["avg_score"], reverse=True)

    monthly_summary = []
    for m in range(1, 13):
        if monthly_stats[m]["total"] > 0:
            monthly_summary.append({
                "month": m,
                "total": monthly_stats[m]["total"],
                "closed": monthly_stats[m]["closed"],
                "plans": monthly_stats[m]["plans"],
                "evidence": monthly_stats[m]["evidence"],
                "avg_score": round(monthly_stats[m]["score"] / monthly_stats[m]["total"], 2) if monthly_stats[m]["total"] > 0 else 0,
            })

    return {
        "year": year,
        "month": month,
        "total_users": total_users,
        "total_indicators": total_indicators,
        "total_tracked": total_tracked,
        "total_closed": total_closed,
        "total_action_plans": total_action_plans,
        "total_evidence": total_evidence,
        "monthly_summary": monthly_summary,
        "teams": teams_list,
    }


def empty_global_dashboard(year: int, month: int = None):
    return {
        "year": year,
        "month": month,
        "total_users": 0,
        "total_indicators": 0,
        "total_tracked": 0,
        "total_closed": 0,
        "total_action_plans": 0,
        "total_evidence": 0,
        "monthly_summary": [],
        "teams": [],
    }


def generate_global_report(db: Session, year: int, area: str = None):
    from app.modules.users.service import normalize_area
    from io import BytesIO
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, Border, Side, PatternFill

    all_users = db.query(User).filter(User.is_active == True).all()

    if area:
        normalized_area = normalize_area(area)
        if normalized_area:
            all_users = [
                u for u in all_users
                if normalize_area(u.area) == normalized_area
            ]

    months_names = ["Enero", "Febrero", "Marzo", "Abril", "Mayo", "Junio",
                    "Julio", "Agosto", "Septiembre", "Octubre", "Noviembre", "Diciembre"]

    headers = [
        "Vicepresidencia", "Área", "Dirección", "Linea", "# Linea", "Cargo",
        "Responsable", "Nombre del Indicador", "Meta", "Peso", "Frecuencia"
    ]
    for m in months_names:
        headers.append(m)
        headers.append(f"Logro {m}")
    headers.append("Observaciones")
    headers.append("Correo Corporativo")

    wb = Workbook()
    ws = wb.active
    ws.title = f"Reporte {year}"

    hfont = Font(bold=True, color="FFFFFF", size=10)
    hfill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    bdr = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin')
    )

    for ci, h in enumerate(headers, 1):
        c = ws.cell(row=1, column=ci, value=h)
        c.font = hfont
        c.fill = hfill
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        c.border = bdr

    rn = 2

    for user in all_users:
        assignments = db.query(IndicatorAssignment).filter(
            IndicatorAssignment.user_id == user.id,
            IndicatorAssignment.year == year,
            IndicatorAssignment.is_active == True
        ).all()

        if not assignments:
            continue

        indicator_names = list(set(a.indicator_name for a in assignments if a.indicator_name))

        for ind_name in indicator_names:
            ind_assignments = [a for a in assignments if a.indicator_name == ind_name]
            first_assignment = ind_assignments[0]

            ind_ids = [a.id for a in ind_assignments]
            trackings = db.query(IndicatorTracking).filter(
                IndicatorTracking.assignment_id.in_(ind_ids)
            ).all()
            trackings_by_month = {t.month: t for t in trackings}

            row_vals = [
                user.area or "",
                user.subarea or "",
                user.direccion or "",
                user.linea or "",
                user.numero_linea or "",
                user.position_name or "",
                user.name,
                ind_name,
                float(first_assignment.target_value) if first_assignment.target_value else "",
                float(first_assignment.weight) if first_assignment.weight else "",
                first_assignment.frequency or "MONTHLY",
            ]

            for m in range(1, 13):
                month_assignment = next((a for a in ind_assignments if a.month == m), None)

                if not month_assignment:
                    row_vals.append("")
                    row_vals.append("X")
                    continue

                meta = float(month_assignment.target_value) if month_assignment and month_assignment.target_value else None
                peso = float(first_assignment.weight) if first_assignment.weight else None

                tracking = trackings_by_month.get(m)
                calificacion = float(tracking.achievement_percentage) if tracking and tracking.achievement_percentage is not None else None

                row_vals.append(calificacion if calificacion is not None else "")

                if calificacion is None:
                    row_vals.append("")
                    continue

                row_vals.append(peso if meta is not None and calificacion >= meta else 0)

            row_vals.append("")
            row_vals.append(user.email or "")

            for ci, val in enumerate(row_vals, 1):
                c = ws.cell(row=rn, column=ci, value=val)
                c.border = bdr
                c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

            rn += 1

    for col_cells in ws.columns:
        max_len = 0
        col_letter = col_cells[0].column_letter
        for cell in col_cells:
            if cell.value:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = min(max_len + 3, 30)

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output